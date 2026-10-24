import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# Sheet 1: Danh sách sản phẩm
ws = wb.active
ws.title = "Danh sách sản phẩm"
ws.views.sheetView[0].showGridLines = True

headers = [
    "ID",
    "SKU (*)",
    "Tên sản phẩm (*)",
    "Mã ngành hàng (*)",
    "Tên ngành hàng",
    "Thương hiệu",
    "Giá bán (VNĐ) (*)",
    "Tồn kho",
    "Trạng thái (1=Hiện, 0=Ẩn)",
    "Đánh giá trung bình",
    "Số lượt đánh giá",
    "Thông số kỹ thuật (JSON/Text)",
    "Danh sách URL ảnh (,)",
    "Đường dẫn sản phẩm gốc"
]

header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
thin_border = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="thin", color="D1D5DB")
)
center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
right_align = Alignment(horizontal="right", vertical="center")

ws.row_dimensions[1].height = 28
for col_idx, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_idx, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center_align
    cell.border = thin_border

data_rows = [
    # Dòng 2: Laptop Acer Predator Helios Neo (Thêm mới)
    [
        "",
        "LAP-ACER-PRED-16",
        "Laptop Acer Predator Helios Neo 16 PHN16-71-779X (i7 13700HX, 16GB, 512GB, RTX 4060, WQXGA 165Hz)",
        "laptop",
        "Laptop",
        "Acer",
        37990000,
        25,
        1,
        4.9,
        34,
        json.dumps({"cpu": "Intel Core i7-13700HX", "ram": 16, "storage": 512, "vga": "RTX 4060 8GB", "screen_size": 16.0, "refresh_rate": 165, "battery": 90}),
        "https://images.unsplash.com/photo-1603302576837-37561b2e2302?w=800, https://images.unsplash.com/photo-1588872657578-7efd1f1555ed?w=800",
        "https://example.com/products/acer-predator-helios-neo-16"
    ],
    # Dòng 3: Laptop Asus Zenbook 14 OLED (Thêm mới)
    [
        "",
        "LAP-ASUS-ZEN14-UX3405",
        "Laptop Asus Zenbook 14 OLED UX3405MA (Core Ultra 7 155H, 32GB, 1TB SSD, 14 inch 3K OLED 120Hz)",
        "laptop",
        "Laptop",
        "Asus",
        32990000,
        18,
        1,
        4.8,
        22,
        json.dumps({"cpu": "Intel Core Ultra 7 155H", "ram": 32, "storage": 1024, "screen_size": 14.0, "refresh_rate": 120, "battery": 75, "weight": 1.28}),
        "https://images.unsplash.com/photo-1496181133206-80ce9b88a853?w=800",
        "https://example.com/products/asus-zenbook-14-oled"
    ],
    # Dòng 4: Phone iPhone 16 Pro Max (Thêm mới ngành phone)
    [
        "",
        "PHO-APPLE-IP16PM-256",
        "iPhone 16 Pro Max 256GB Titan Sa Mạc",
        "phone",
        "Điện thoại",
        "Apple",
        34990000,
        40,
        1,
        5.0,
        95,
        json.dumps({"chipset": "Apple A18 Pro", "ram": 8, "storage": 256, "battery": 4685, "screen_size": 6.9, "camera": 48}),
        "https://images.unsplash.com/photo-1510557880182-3d4d3cba35a5?w=800",
        "https://example.com/products/iphone-16-pro-max-256gb"
    ],
    # Dòng 5: Phone Samsung Galaxy S24 Ultra (Thêm mới ngành phone)
    [
        "",
        "PHO-SAM-S24U-512",
        "Samsung Galaxy S24 Ultra 5G 12GB/512GB Titan Xám",
        "phone",
        "Điện thoại",
        "Samsung",
        31990000,
        30,
        1,
        4.8,
        68,
        json.dumps({"chipset": "Snapdragon 8 Gen 3 for Galaxy", "ram": 12, "storage": 512, "battery": 5000, "screen_size": 6.8, "camera": 200}),
        "https://images.unsplash.com/photo-1610945415295-d9bbf067e59c?w=800",
        "https://example.com/products/samsung-galaxy-s24-ultra"
    ],
    # Dòng 6: Phone Xiaomi 14 Ultra (Thêm mới)
    [
        "",
        "PHO-XIAOMI-14U",
        "Xiaomi 14 Ultra 16GB/512GB Leica Camera",
        "phone",
        "Điện thoại",
        "Xiaomi",
        27990000,
        15,
        1,
        4.7,
        41,
        json.dumps({"chipset": "Snapdragon 8 Gen 3", "ram": 16, "storage": 512, "battery": 5000, "screen_size": 6.73, "camera": 50}),
        "https://images.unsplash.com/photo-1598327105666-5b89351aff97?w=800",
        "https://example.com/products/xiaomi-14-ultra"
    ],
    # Dòng 7: Dòng cập nhật Upsert (Cập nhật cho sản phẩm có SKU LAP-ASUS-VIVO15 trong mẫu template)
    [
        "",
        "LAP-ASUS-VIVO15",
        "Laptop Asus Vivobook 15 OLED A1505ZA [Bản Cập Nhật Giá Mới]",
        "laptop",
        "Laptop",
        "Asus",
        15490000,
        50,
        1,
        4.9,
        30,
        json.dumps({"ram": 16, "storage": 512, "cpu": "Core i5-1235U", "screen_size": 15.6}),
        "https://images.unsplash.com/photo-1496181133206-80ce9b88a853?w=800",
        "https://example.com/products/asus-vivobook-15"
    ],
    # Dòng 8: Dòng thiếu SKU để test cơ chế tự sinh SKU của hệ thống
    [
        "",
        "",
        "Laptop Lenovo IdeaPad Slim 5 14IRH9 (Core i5 13500H, 16GB, 512GB, OLED 120Hz)",
        "laptop",
        "Laptop",
        "Lenovo",
        18990000,
        12,
        1,
        4.6,
        15,
        json.dumps({"cpu": "Intel Core i5-13500H", "ram": 16, "storage": 512, "screen_size": 14.0, "battery": 57}),
        "https://images.unsplash.com/photo-1541807084-5c52b6b3adef?w=800",
        "https://example.com/products/lenovo-ideapad-slim-5"
    ]
]

for row_idx, r_data in enumerate(data_rows, 2):
    ws.row_dimensions[row_idx].height = 24
    for col_idx, val in enumerate(r_data, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.border = thin_border
        if col_idx in [1, 2, 4, 8, 9, 11]:  # ID, SKU, Code, Stock, Status, ReviewCount
            cell.alignment = center_align
        elif col_idx == 7:  # Price
            cell.alignment = right_align
            cell.number_format = '#,##0'
        elif col_idx == 10:  # Rating
            cell.alignment = center_align
            cell.number_format = '0.0'
        else:
            cell.alignment = left_align

# Sheet 2: Hướng dẫn kịch bản kiểm thử
ws_guide = wb.create_sheet(title="Hướng dẫn test Import")
ws_guide.views.sheetView[0].showGridLines = True
ws_guide.row_dimensions[1].height = 26

guide_headers = ["Dòng Excel", "Loại Test Case", "Mô tả & Kỳ vọng hệ thống"]
for col_idx, h in enumerate(guide_headers, 1):
    cell = ws_guide.cell(row=1, column=col_idx, value=h)
    cell.font = header_font
    cell.fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    cell.alignment = center_align
    cell.border = thin_border

guide_data = [
    ["Dòng 2 - 4", "Thêm mới sản phẩm Laptop & Phone", "Thêm mới hoàn toàn 3 sản phẩm có đầy đủ ảnh, specs JSON, SKU chuẩn. Kỳ vọng: 'createdCount' tăng."],
    ["Dòng 5 - 6", "Thêm mới dòng Flagship (Samsung, Xiaomi)", "Kiểm tra độ chính xác của trường số lớn (camera 200MP, giá 31.990.000đ). Kỳ vọng: 'createdCount' tăng."],
    ["Dòng 7", "Thử nghiệm tính năng UPSERT (Cập nhật)", "Dùng SKU 'LAP-ASUS-VIVO15'. Nếu đã có trong CSDL, hệ thống cập nhật giá mới 15.490.000đ và tồn kho 50. Kỳ vọng: 'updatedCount' tăng."],
    ["Dòng 8", "Tự động sinh mã SKU", "Cố ý để trống cột SKU. Hệ thống tự động sinh SKU có dạng SKU-LAPTOP-<timestamp>-8 và ghi Warning. Kỳ vọng: 'createdCount' tăng kèm warning log, không bị lỗi."]
]

for row_idx, r_data in enumerate(guide_data, 2):
    ws_guide.row_dimensions[row_idx].height = 24
    for col_idx, val in enumerate(r_data, 1):
        cell = ws_guide.cell(row=row_idx, column=col_idx, value=val)
        cell.border = thin_border
        cell.alignment = left_align if col_idx == 3 else center_align

# Tự căn chỉnh chiều rộng cột
for sheet in [ws, ws_guide]:
    for col in sheet.columns:
        max_len = 0
        for cell in col:
            v = str(cell.value or '')
            if len(v) > max_len:
                max_len = len(v)
        col_letter = get_column_letter(col[0].column)
        sheet.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 48)

file_path = "d:/AI_Shopping_Agent/product_import_test_data.xlsx"
wb.save(file_path)
print(f"File created successfully: {file_path}")
